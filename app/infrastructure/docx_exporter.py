"""Adaptador del puerto DocumentExporter: Markdown a Word con markdown-it-py y python-docx.

El Markdown se interpreta una sola vez con un parser real (CommonMark) y el
documento se construye recorriendo sus tokens.
"""

import io

from docx import Document
from docx.document import Document as DocxDocument
from docx.shared import RGBColor
from docx.text.paragraph import Paragraph
from markdown_it import MarkdownIt
from markdown_it.token import Token

from app.application.interfaces.document_exporter import DocumentExporter

TITLE_COLOR = RGBColor(0x1A, 0x56, 0xDB)
CODE_FONT = "Consolas"
LIST_STYLES = {"bullet": "List Bullet", "ordered": "List Number"}
# La plantilla por defecto de Word trae estilos de lista hasta el nivel 3
MAX_LIST_LEVEL = 3
MAX_HEADING_LEVEL = 9


class DocxExporter(DocumentExporter):
    def __init__(self):
        self._markdown = MarkdownIt("commonmark")

    def export(self, title: str, markdown: str) -> bytes:
        document = Document()
        heading = document.add_heading(level=1)
        heading.add_run(title).font.color.rgb = TITLE_COLOR

        self._render(document, self._markdown.parse(markdown))

        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    def _render(self, document: DocxDocument, tokens: list[Token]) -> None:
        open_lists: list[str] = []
        heading_level: int | None = None

        for token in tokens:
            match token.type:
                case "heading_open":
                    heading_level = int(token.tag[1])
                case "heading_close":
                    heading_level = None
                case "bullet_list_open":
                    open_lists.append("bullet")
                case "ordered_list_open":
                    open_lists.append("ordered")
                case "bullet_list_close" | "ordered_list_close":
                    open_lists.pop()
                case "inline":
                    paragraph = self._new_paragraph(document, heading_level, open_lists)
                    self._add_inline(paragraph, token.children or [])
                case "fence" | "code_block":
                    paragraph = document.add_paragraph()
                    paragraph.add_run(token.content.rstrip("\n")).font.name = CODE_FONT

    @staticmethod
    def _new_paragraph(
        document: DocxDocument, heading_level: int | None, open_lists: list[str]
    ) -> Paragraph:
        if heading_level is not None:
            # El nivel 1 es el título del documento: los del Markdown van debajo
            return document.add_heading(level=min(heading_level + 1, MAX_HEADING_LEVEL))
        if open_lists:
            depth = min(len(open_lists), MAX_LIST_LEVEL)
            style = LIST_STYLES[open_lists[-1]] + ("" if depth == 1 else f" {depth}")
            return document.add_paragraph(style=style)
        return document.add_paragraph()

    @staticmethod
    def _add_inline(paragraph: Paragraph, children: list[Token]) -> None:
        bold = italic = False

        for child in children:
            match child.type:
                case "strong_open" | "strong_close":
                    bold = child.type == "strong_open"
                case "em_open" | "em_close":
                    italic = child.type == "em_open"
                case "text":
                    run = paragraph.add_run(child.content)
                    run.bold = bold or None
                    run.italic = italic or None
                case "code_inline":
                    paragraph.add_run(child.content).font.name = CODE_FONT
                case "softbreak":
                    paragraph.add_run(" ")
                case "hardbreak":
                    paragraph.add_run().add_break()
